"""Build an editable PowerPoint and copy-ready speaking guide from slide content.

Optional authoring dependency: uv run --with python-pptx==1.0.2 python -m scripts.build_presentation
This dependency is not required by the application.
"""

import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/presentation/slide-content.json"
OUTPUT = ROOT / "docs/presentation/Jinni_Code_Walkthrough.pptx"
GUIDE = ROOT / "docs/PRESENTATION_GUIDE.md"
COLORS = ("BC3045", "18704C", "255EC9", "8A630A")
SOFT = ("FFF0EE", "EAF7EF", "EDF3FF", "FFF7DB")
INK, MUTED, BACKGROUND = "17243B", "526078", "FAFBFE"


def verify_worked_example():
    """Prevent a changed fixture/ranking from silently leaving the slides inaccurate."""
    from fastapi.testclient import TestClient

    from app.main import create_app

    with TestClient(create_app()) as client:
        entry = client.get("/api/v1/decision-scenarios/travel").json()["data"]
        result = client.post(
            "/api/v1/decisions/evaluate", json=entry["scenario"]["default_request"]
        ).json()["data"]
    expected = [("train", 55.0, 1800, 70, 2200), ("coach", 45.0, 800, 120, 5000)]
    actual = [
        (
            item["id"],
            item["score"],
            *[
                item["metrics"][key]
                for key in ("cost-pence", "duration-minutes", "emissions-grams")
            ],
        )
        for item in result["ranked"]
    ]
    if actual != expected:
        raise ValueError("The travel example changed. Update and verify slide 8 before rebuilding.")


def write_guide(content):
    lines = [
        "# Jinni: what to write in your PowerPoint",
        "",
        "The editable deck is "
        "[Jinni_Code_Walkthrough.pptx](presentation/Jinni_Code_Walkthrough.pptx). "
        "This guide contains the slide text, what to say, and files to explain. "
        "Speaker notes are also embedded in the deck.",
        "",
        "Aim for 10–12 minutes. Add your name and course/event details on the title slide. "
        "Use the slide text as prompts; explain the speaker notes in your own words. "
        "Keep screenshots and demo examples visibly labelled fictional. "
        "No genuine API key or live OpenAI success is included in this deck.",
        "",
        "Edit [slide-content.json](presentation/slide-content.json) to change the source material. "
        "Regenerate with `uv run --with python-pptx==1.0.2 python -m scripts.build_presentation`. "
        "The optional authoring dependency does not change application requirements.",
        "",
    ]
    for index, data in enumerate(content["slides"], 1):
        lines += [f"## Slide {index}: {data['title']}", "", "**Write on the slide**", ""]
        if data.get("subtitle"):
            lines += [data["subtitle"], ""]
        if data.get("formula"):
            lines += [f"`{data['formula']}`", ""]
        if data.get("headers"):
            lines += [
                "| " + " | ".join(data["headers"]) + " |",
                "| " + " | ".join("---" for _ in data["headers"]) + " |",
            ]
        if data.get("rows"):
            if not data.get("headers"):
                lines += ["| Code | Responsibility |", "| --- | --- |"]
            lines += ["| " + " | ".join(row) + " |" for row in data["rows"]]
            lines.append("")
        items = list(data.get("bullets", []))
        items += [
            f"{card['title']}: {card['body'].replace(chr(10), ' ')}"
            for card in data.get("cards", [])
        ]
        items += [f"{step['title']}: {step['body']}" for step in data.get("steps", [])]
        for side in ("left", "right"):
            if side in data:
                lines += [f"**{data[side]['title']}**", ""]
                lines += [f"- {text}" for text in data[side]["bullets"]] + [""]
        lines += [f"- {text}" for text in items]
        if data.get("callout"):
            lines += ["", data["callout"]]
        lines += ["", "**What to say**", "", data["notes"], "", "**Code to explain**", ""]
        lines += [f"- `{path}`" for path in data["references"]]
        lines.append("")
    lines += [
        "## Questions you should be ready to answer",
        "",
        "- **Why use Python and an AI model?** Python applies repeatable constraints and scoring; "
        "OpenAI explains the computed result using selected evidence.",
        "- **What does a score of 100 mean?** Best fit under the supplied option set and settings. "
        "It is not a 100% chance of success, and a single feasible option can score 100.",
        "- **Why integer pence?** £7.99 and £8.00 remain exact budget boundaries.",
        "- **Are citations a truth guarantee?** They establish that the quote exists in "
        "the supplied "
        "evidence. They do not prove every sentence follows from that evidence.",
        "- **Does the website use OpenAI?** The GitHub Pages edition uses local calculations and "
        "text search. The desktop edition has the provider-backed AI tools.",
        "- **Can it analyse any spreadsheet?** Current imports are bounded "
        "text/Markdown/text PDFs. "
        "Custom structured comparisons use the declared criteria/data schema; there is no CSV, "
        "Excel, OCR, forecasting or arbitrary-code analysis engine.",
        "",
        "See [OPENAI_DESKTOP_GUIDE.md](OPENAI_DESKTOP_GUIDE.md) for setup, prompts, "
        "the demo checklist and the optional live verification command.",
        "",
    ]
    GUIDE.write_text("\n".join(lines), encoding="utf-8")


def main():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    content = json.loads(SOURCE.read_text(encoding="utf-8"))
    verify_worked_example()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = content["title"]
    deck.core_properties.subject = "Code documentation, OpenAI setup and evidence boundaries"
    deck.core_properties.author = "Jinni project"
    deck.core_properties.keywords = "Jinni, code walkthrough, OpenAI, fictional data"

    def box(slide, x, y, width, height, color, rounded=True):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        if rounded:
            shape.adjustments[0] = 0.08
        return shape

    def text(slide, x, y, width, height, value, size=20, color=INK, bold=False, mono=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.01)
        frame.margin_top = frame.margin_bottom = 0
        lines = value.split("\n")
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Consolas" if mono else "Aptos"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(8 if index < len(lines) - 1 else 0)
        return shape

    for index, data in enumerate(content["slides"], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BACKGROUND)
        for segment, color in enumerate(COLORS):
            box(slide, 0.65 + segment * 0.25, 0.4, 0.25, 0.055, color, False)
        text(slide, 0.65, 0.62, 11.9, 0.28, data["kicker"], 11, COLORS[2], True)
        text(slide, 0.65, 1.02, 12.0, 0.75, data["title"], 34, bold=True)
        text(
            slide, 0.65, 7.06, 10.8, 0.2, "JINNI · CODE WALKTHROUGH · FICTIONAL DEMO DATA", 9, MUTED
        )
        text(slide, 12.1, 7.02, 0.65, 0.25, f"{index:02} / {len(content['slides'])}", 10, MUTED)
        layout = data["layout"]
        if layout == "cover":
            text(slide, 0.65, 2.35, 11.95, 1.5, data["subtitle"], 44, bold=True)
            for item, label in enumerate(data["bullets"]):
                box(slide, 0.65 + item * 4.05, 4.8, 3.85, 1.15, SOFT[item])
                text(slide, 0.88 + item * 4.05, 5.02, 3.4, 0.8, label, 22, COLORS[item], True)
            text(
                slide,
                0.65,
                6.32,
                11.8,
                0.35,
                "Desktop application + optional OpenAI · presentation notes included",
                16,
                MUTED,
            )
        elif layout in ("cards", "process"):
            items = data.get("cards", data.get("steps", []))
            width = (12.03 - (len(items) - 1) * 0.22) / len(items)
            for item, entry in enumerate(items):
                x = 0.65 + item * (width + 0.22)
                box(slide, x, 2.2, width, 3.25, SOFT[item % 4])
                text(
                    slide,
                    x + 0.22,
                    2.4,
                    width - 0.44,
                    0.35,
                    f"{item + 1:02}",
                    14,
                    COLORS[item % 4],
                    True,
                )
                text(
                    slide,
                    x + 0.22,
                    3.0,
                    width - 0.44,
                    0.65,
                    entry["title"],
                    23,
                    COLORS[item % 4],
                    True,
                )
                text(slide, x + 0.22, 3.85, width - 0.44, 1.4, entry["body"], 20)
        elif layout == "split":
            for item, side in enumerate(("left", "right")):
                x = 0.65 + item * 6.12
                box(slide, x, 1.98, 5.9, 4.22, SOFT[item + 1])
                text(
                    slide,
                    x + 0.25,
                    2.2,
                    5.38,
                    0.45,
                    data[side]["title"],
                    23,
                    COLORS[item + 1],
                    True,
                )
                for line, value in enumerate(data[side]["bullets"]):
                    text(slide, x + 0.25, 2.95 + line * 0.73, 5.35, 0.72, "• " + value, 18)
        elif layout in ("map", "example"):
            headers = data.get("headers", ["CODE", "RESPONSIBILITY"])
            rows = data["rows"]
            widths = [4.5, 7.53] if layout == "map" else [3.2, 1.6, 2.5, 2.4, 2.33]
            y = 1.98
            for row, values in enumerate([headers, *rows]):
                height = 0.55 if layout == "map" else 0.64
                x = 0.65
                for column, value in enumerate(values):
                    box(
                        slide,
                        x,
                        y,
                        widths[column],
                        height,
                        "E7EDFA" if row == 0 else ("FFFFFF" if row % 2 else "F0F4FB"),
                        False,
                    )
                    text(
                        slide,
                        x + 0.12,
                        y + 0.11,
                        widths[column] - 0.24,
                        height - 0.12,
                        value,
                        13 if layout == "map" and column == 0 else 17,
                        COLORS[2] if row == 0 else INK,
                        row == 0,
                        layout == "map" and column == 0,
                    )
                    x += widths[column]
                y += height + 0.04
            for row, value in enumerate(data.get("bullets", [])):
                text(slide, 0.7, 4.55 + row * 0.48, 11.8, 0.45, "• " + value, 18)
        elif layout == "formula":
            box(slide, 0.65, 2.05, 12.03, 1.0, SOFT[2])
            text(slide, 0.9, 2.32, 11.5, 0.6, data["formula"], 23, COLORS[2], True)
            for item, value in enumerate(data["bullets"]):
                text(slide, 0.85, 3.55 + item * 0.55, 11.7, 0.5, f"{item + 1}.  {value}", 22)
        elif layout == "checklist":
            for item, value in enumerate(data["bullets"]):
                y = 2.0 + item * 0.7
                box(slide, 0.65, y, 0.45, 0.45, SOFT[item % 4])
                text(slide, 0.76, y + 0.06, 0.25, 0.3, str(item + 1), 16, COLORS[item % 4], True)
                text(slide, 1.3, y + 0.02, 11.1, 0.6, value, 20)
            box(slide, 0.65, 5.9, 12.03, 0.82, SOFT[2])
            text(slide, 0.9, 6.08, 11.45, 0.58, data["callout"], 17, COLORS[2], True)
        else:
            raise ValueError(f"Unknown slide layout: {layout}")
        if data.get("subtitle") and layout != "cover":
            text(slide, 0.7, 6.38, 11.95, 0.45, data["subtitle"], 16, MUTED)
        slide.notes_slide.notes_text_frame.text = (
            data["notes"] + "\n\nCode references:\n" + "\n".join(data["references"])
        )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    write_guide(content)
    reopened = Presentation(OUTPUT)
    if len(reopened.slides) != len(content["slides"]):
        raise ValueError("The saved presentation is missing slides")
    for slide in reopened.slides:
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError("A slide is missing its speaking notes")
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > deck.slide_width
                or shape.top + shape.height > deck.slide_height
            ):
                raise ValueError("A slide shape extends outside the slide")
    print(f"Built {len(reopened.slides)} editable slides with speaker notes: {OUTPUT}")
    print(f"Copy-ready writing guide: {GUIDE}")


if __name__ == "__main__":
    main()
